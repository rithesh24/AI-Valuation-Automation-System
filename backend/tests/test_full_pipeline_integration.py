"""Integration test: proves the pieces already unit-tested in isolation
(upload, document_parser, claude_service, report_service) actually compose
into one working pipeline. Only the outermost Claude API boundary
(ClaudeService._call_claude) is faked; uploads, parsing, the mapping cache,
and DOCX injection all run for real. See docs/ROADMAP.md Phase 8.
"""

from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services.claude_service import ClaudeService
from app.services.document_parser import DocumentParser

client = TestClient(app)


class _TextBlock:
    def __init__(self, text: str) -> None:
        self.type = "text"
        self.text = text


class _Usage:
    def __init__(self) -> None:
        self.input_tokens = 100
        self.output_tokens = 50


class _FakeResponse:
    def __init__(self, text: str) -> None:
        self.content = [_TextBlock(text)]
        self.usage = _Usage()


@pytest.fixture(autouse=True)
def _isolated_storage(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("app.core.config.settings.UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setattr("app.core.config.settings.REPORTS_DIR", str(tmp_path / "reports"))
    monkeypatch.setattr("app.core.db.settings.DATABASE_PATH", str(tmp_path / "test.db"))
    monkeypatch.setattr("app.services.claude_service.settings.ANTHROPIC_API_KEY", "sk-test")


@pytest.fixture(autouse=True)
def _fake_claude_api_boundary(monkeypatch) -> None:
    def fake_call_claude(self, prompt, tools, on_status=None):  # noqa: ARG001
        if tools:  # extraction attaches the web-search tool (D8 Tier 2); mapping doesn't
            return _FakeResponse('{"property_identification": {"district": "Pune"}}')
        return _FakeResponse(
            '{"property_identification.district": {"location_id": "p1", "confidence": "high"}}'
        )

    monkeypatch.setattr(ClaudeService, "_call_claude", fake_call_claude)


def _build_property_document(path: Path) -> None:
    document = Document()
    document.add_paragraph("Property in Pune district.")
    document.save(path)


def _build_template(path: Path) -> None:
    document = Document()
    document.add_paragraph("VALUATION REPORT")
    document.add_paragraph("")  # district value line -> location_id "p1"
    document.save(path)


def test_upload_parse_extract_map_inject_pipeline(tmp_path: Path) -> None:
    doc_path = tmp_path / "property.docx"
    _build_property_document(doc_path)

    with doc_path.open("rb") as f:
        upload_response = client.post(
            "/uploads",
            files={
                "files": (
                    "property.docx",
                    f,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            data={"category": "property_document"},
        )
    assert upload_response.status_code == 200
    uploaded_path = upload_response.json()[0]["path"]

    document_text = DocumentParser().extract_docx_text(uploaded_path)
    assert "Pune" in document_text

    extraction = ClaudeService().extract_valuation_data(
        document_texts={"property.docx": document_text}
    )
    assert extraction.data.property_identification.district == "Pune"

    template_path = tmp_path / "template.docx"
    _build_template(template_path)

    generate_response = client.post(
        "/reports/generate",
        json={"template_path": str(template_path), "data": extraction.data.model_dump()},
    )
    assert generate_response.status_code == 200
    report_id = generate_response.json()["report_id"]

    preview = client.get(f"/reports/{report_id}/preview")
    assert "Pune" in preview.json()["text"]
