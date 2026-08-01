from datetime import date, datetime, timezone
from pathlib import Path

import fitz
import pytest
from docx import Document
from docx.shared import Inches

from app.models.template_mapping import FieldMapping
from app.models.valuation_schema import (
    INFO_NOT_AVAILABLE,
    ComparableEvidence,
    ValuationReportData,
    canonical_schema_field_list,
)
from app.services.claude_service import ClaudeService
from app.services.easr_service import EASRGuidelineResult, EASRSearchInput, EASRService
from app.services.report_service import (
    ReportService,
    ReportServiceError,
    _citation_from_easr_result,
    _financial_year_string,
    _resolve_template_path,
)


def _build_template(path: Path) -> None:
    document = Document()
    document.add_paragraph("VALUATION REPORT")
    document.add_paragraph("Name of Owner")
    document.add_paragraph("")  # blank value line under the label above
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "District"
    table.rows[0].cells[1].text = ""
    table.rows[1].cells[0].text = "Taluka"
    table.rows[1].cells[1].text = ""
    document.save(path)


class TestExtractTemplateSkeleton:
    def test_extracts_paragraphs_in_document_order(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_template(path)

        skeleton = ReportService().extract_template_skeleton(str(path))
        paragraph_texts = [loc.text for loc in skeleton if loc.type == "paragraph"]

        assert paragraph_texts == ["VALUATION REPORT", "Name of Owner", ""]

    def test_blank_paragraph_context_is_previous_non_blank_paragraph(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_template(path)

        skeleton = ReportService().extract_template_skeleton(str(path))
        blank_value_line = next(loc for loc in skeleton if loc.type == "paragraph" and loc.text == "")

        assert blank_value_line.context == ["Name of Owner"]

    def test_table_cell_context_is_the_rest_of_its_row(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_template(path)

        skeleton = ReportService().extract_template_skeleton(str(path))
        district_label = next(loc for loc in skeleton if loc.text == "District")
        district_value = next(
            loc for loc in skeleton if loc.type == "table_cell" and loc.location_id == "t0r0c1"
        )

        assert district_label.context == [""]
        assert district_value.context == ["District"]

    def test_location_ids_are_unique(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_template(path)

        skeleton = ReportService().extract_template_skeleton(str(path))
        ids = [loc.location_id for loc in skeleton]

        assert len(ids) == len(set(ids))

    def test_missing_file_raises_report_service_error(self) -> None:
        try:
            ReportService().extract_template_skeleton("does_not_exist.docx")
            assert False, "expected ReportServiceError"
        except ReportServiceError:
            pass


class TestSkeletonHash:
    def test_stable_for_identical_content(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_template(path)
        service = ReportService()

        skeleton_a = service.extract_template_skeleton(str(path))
        skeleton_b = service.extract_template_skeleton(str(path))

        assert service.skeleton_hash(skeleton_a) == service.skeleton_hash(skeleton_b)

    def test_stable_when_only_text_changes(self, tmp_path: Path) -> None:
        """Real client templates are previous filled reports, not blank forms (D28) —
        two uploads of the same bank format with different property data (different
        owner name, different district value, even a relabelled field) must still
        hit the same cache entry, since the hash is structural only."""
        path = tmp_path / "template.docx"
        _build_template(path)
        service = ReportService()
        original_hash = service.skeleton_hash(service.extract_template_skeleton(str(path)))

        document = Document(str(path))
        document.paragraphs[1].text = "Name of Purchaser"  # label
        document.paragraphs[2].text = "Mr. Bhavesh Ishwarbhai Darji"  # value
        document.tables[0].rows[0].cells[1].text = "Mumbai Suburban"  # value
        document.save(path)

        changed_hash = service.skeleton_hash(service.extract_template_skeleton(str(path)))

        assert original_hash == changed_hash

    def test_changes_when_a_table_column_is_added(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_template(path)
        service = ReportService()
        original_hash = service.skeleton_hash(service.extract_template_skeleton(str(path)))

        document = Document(str(path))
        document.tables[0].add_column(Inches(1))
        document.save(path)

        changed_hash = service.skeleton_hash(service.extract_template_skeleton(str(path)))

        assert original_hash != changed_hash


class _FakeClaudeService:
    """Stands in for ClaudeService in mapping-cache tests — no real API call."""

    def __init__(self, mapping: dict[str, FieldMapping]) -> None:
        self.mapping = mapping
        self.calls = 0

    def map_template_fields(self, _skeleton) -> dict[str, FieldMapping]:
        self.calls += 1
        return self.mapping


@pytest.fixture(autouse=True)
def _isolated_database(tmp_path: Path, monkeypatch) -> None:
    """Points the SQLite cache at a fresh per-test file so tests don't share state."""
    monkeypatch.setattr(
        "app.core.db.settings.DATABASE_PATH", str(tmp_path / "test.db")
    )


class TestMappingCache:
    def test_cache_miss_calls_claude_and_saves(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_template(path)
        fake_mapping = {"property_identification.district": FieldMapping(location_id="t0r0c1")}
        claude = _FakeClaudeService(fake_mapping)
        service = ReportService()

        mapping, skeleton_hash = service.get_or_create_mapping(str(path), claude_service=claude)

        assert claude.calls == 1
        assert mapping["property_identification.district"] == fake_mapping["property_identification.district"]
        # get_or_create_mapping backfills every canonical field Claude's response
        # didn't return (see _normalize_mapping) so it — and the cached copy —
        # are always complete, not just whatever the fake happened to include.
        assert len(mapping) == len(canonical_schema_field_list())
        assert service.get_cached_mapping(skeleton_hash) == mapping

    def test_cache_hit_does_not_call_claude(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_template(path)
        fake_mapping = {"property_identification.district": FieldMapping(location_id="t0r0c1")}
        claude = _FakeClaudeService(fake_mapping)
        service = ReportService()
        service.get_or_create_mapping(str(path), claude_service=claude)

        mapping, _ = service.get_or_create_mapping(str(path), claude_service=claude)

        assert claude.calls == 1  # not called a second time
        assert mapping["property_identification.district"] == fake_mapping["property_identification.district"]

    def test_force_regenerate_calls_claude_again(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_template(path)
        claude = _FakeClaudeService(
            {"property_identification.district": FieldMapping(location_id="t0r0c1")}
        )
        service = ReportService()
        service.get_or_create_mapping(str(path), claude_service=claude)

        service.get_or_create_mapping(str(path), claude_service=claude, force_regenerate=True)

        assert claude.calls == 2


def _build_pdf_template(path: Path) -> None:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "VALUATION REPORT")
    page.insert_text((72, 110), "Name of Owner")
    document.save(str(path))
    document.close()


class TestPdfTemplateSupport:
    """A born-digital PDF template is converted to .docx once, then flows
    through the exact same Stage 2/3 pipeline as a native .docx template."""

    def test_pdf_is_converted_and_skeleton_matches_its_text(self, tmp_path: Path) -> None:
        pdf_path = tmp_path / "template.pdf"
        _build_pdf_template(pdf_path)

        resolved = _resolve_template_path(str(pdf_path))
        converted_docx = tmp_path / ".converted" / "template.docx"

        assert resolved == str(converted_docx)
        assert converted_docx.exists()
        skeleton = ReportService().extract_template_skeleton(resolved)
        combined_text = " ".join(loc.text for loc in skeleton)
        assert "VALUATION REPORT" in combined_text
        assert "Name of Owner" in combined_text

    def test_docx_template_is_returned_unchanged(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_template(path)

        assert _resolve_template_path(str(path)) == str(path)

    def test_conversion_is_cached_not_repeated(self, tmp_path: Path, monkeypatch) -> None:
        pdf_path = tmp_path / "template.pdf"
        _build_pdf_template(pdf_path)
        _resolve_template_path(str(pdf_path))

        conversion_calls = []
        import pdf2docx

        original_converter = pdf2docx.Converter

        def _tracking_converter(*args, **kwargs):
            conversion_calls.append(1)
            return original_converter(*args, **kwargs)

        monkeypatch.setattr("pdf2docx.Converter", _tracking_converter)

        _resolve_template_path(str(pdf_path))

        assert conversion_calls == []

    def test_get_or_create_mapping_accepts_a_pdf_template(self, tmp_path: Path) -> None:
        pdf_path = tmp_path / "template.pdf"
        _build_pdf_template(pdf_path)
        fake_mapping = {"ownership_and_title.present_owners": FieldMapping(location_id="p0")}
        claude = _FakeClaudeService(fake_mapping)

        mapping, _ = ReportService().get_or_create_mapping(str(pdf_path), claude_service=claude)

        assert claude.calls == 1
        assert mapping["ownership_and_title.present_owners"] == fake_mapping["ownership_and_title.present_owners"]

    def test_inject_data_fills_a_converted_pdf_template(self, tmp_path: Path) -> None:
        pdf_path = tmp_path / "template.pdf"
        _build_pdf_template(pdf_path)
        service = ReportService()

        resolved = _resolve_template_path(str(pdf_path))
        skeleton = service.extract_template_skeleton(resolved)
        owner_location = next(loc for loc in skeleton if "Name of Owner" in loc.text)
        mapping = {
            "ownership_and_title.present_owners": FieldMapping(location_id=owner_location.location_id)
        }
        data = ValuationReportData()
        data.ownership_and_title.present_owners = "John Smith"
        output_path = tmp_path / "output.docx"

        result = service.inject_data(str(pdf_path), mapping, data, str(output_path))

        assert "ownership_and_title.present_owners" in result.filled_fields
        output_text = "\n".join(p.text for p in Document(str(output_path)).paragraphs)
        assert "John Smith" in output_text

    def test_get_cached_mapping_returns_none_when_absent(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_template(path)
        service = ReportService()
        skeleton_hash = service.skeleton_hash(service.extract_template_skeleton(str(path)))

        assert service.get_cached_mapping(skeleton_hash) is None


def _build_injection_template(path: Path) -> None:
    document = Document()
    document.add_paragraph("VALUATION REPORT")
    paragraph = document.add_paragraph()
    run = paragraph.add_run("PLACEHOLDER")
    run.bold = True
    run.font.name = "Times New Roman"
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Indicator"
    table.rows[0].cells[1].text = "Location"
    document.save(path)


class TestInjectData:
    def test_scalar_field_preserves_run_formatting(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_injection_template(path)
        mapping = {"property_identification.district": FieldMapping(location_id="p1")}
        data = ValuationReportData(property_identification={"district": "Pune"})
        out_path = tmp_path / "out.docx"

        result = ReportService().inject_data(str(path), mapping, data, str(out_path))

        assert result.filled_fields == ["property_identification.district"]
        out_doc = Document(str(out_path))
        run = out_doc.paragraphs[1].runs[0]
        assert run.text == "Pune"
        assert run.bold is True
        assert run.font.name == "Times New Roman"

    def test_repeating_group_clones_rows_and_fills_columns(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_injection_template(path)
        mapping = {
            "comparable_evidence[].indicator_number": FieldMapping(location_id="t0r1c0"),
            "comparable_evidence[].location": FieldMapping(location_id="t0r1c1"),
        }
        data = ValuationReportData(
            comparable_evidence=[
                ComparableEvidence(indicator_number="1", location="Kothrud"),
                ComparableEvidence(indicator_number="2", location="Baner"),
                ComparableEvidence(indicator_number="3", location="Wakad"),
            ]
        )
        out_path = tmp_path / "out.docx"

        ReportService().inject_data(str(path), mapping, data, str(out_path))

        out_doc = Document(str(out_path))
        rows = [[cell.text for cell in row.cells] for row in out_doc.tables[0].rows]
        assert rows == [
            ["Indicator", "Location"],
            ["1", "Kothrud"],
            ["2", "Baner"],
            ["3", "Wakad"],
        ]

    def test_unmapped_field_with_real_value_is_reported(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_injection_template(path)
        mapping = {
            "property_identification.district": FieldMapping(location_id=None, status="not_present")
        }
        data = ValuationReportData(property_identification={"district": "Pune"})
        out_path = tmp_path / "out.docx"

        result = ReportService().inject_data(str(path), mapping, data, str(out_path))

        assert result.unmapped_fields == ["property_identification.district"]

    def test_field_marked_unavailable_is_not_reported_as_unmapped(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_injection_template(path)
        mapping = {
            "property_identification.district": FieldMapping(location_id=None, status="not_present")
        }
        data = ValuationReportData()  # district defaults to the non-fabrication sentinel
        out_path = tmp_path / "out.docx"

        result = ReportService().inject_data(str(path), mapping, data, str(out_path))

        assert result.unmapped_fields == []
        assert result.filled_fields == []

    def test_empty_repeating_list_leaves_prototype_row_untouched(self, tmp_path: Path) -> None:
        path = tmp_path / "template.docx"
        _build_injection_template(path)
        mapping = {"comparable_evidence[].indicator_number": FieldMapping(location_id="t0r1c0")}
        data = ValuationReportData(comparable_evidence=[])
        out_path = tmp_path / "out.docx"

        result = ReportService().inject_data(str(path), mapping, data, str(out_path))

        assert result.filled_fields == []
        out_doc = Document(str(out_path))
        assert len(out_doc.tables[0].rows) == 2  # no rows cloned/removed


class TestRunQualityCheck:
    def test_flags_unmapped_field_with_real_value(self) -> None:
        mapping = {
            "property_identification.district": FieldMapping(location_id=None, status="not_present")
        }
        data = ValuationReportData(property_identification={"district": "Pune"})

        result = ReportService().run_quality_check(mapping, data)

        assert result.passed is False
        assert result.injection_failures == ["property_identification.district"]
        assert result.disclosed_unavailable == []

    def test_disclosed_unavailable_is_not_a_failure(self) -> None:
        mapping = {
            "property_identification.district": FieldMapping(location_id=None, status="not_present")
        }
        data = ValuationReportData()  # sentinel default

        result = ReportService().run_quality_check(mapping, data)

        assert result.passed is True
        assert result.disclosed_unavailable == ["property_identification.district"]
        assert result.injection_failures == []

    def test_mapped_field_with_real_value_passes(self) -> None:
        mapping = {
            "property_identification.district": FieldMapping(location_id="p1", confidence="high")
        }
        data = ValuationReportData(property_identification={"district": "Pune"})

        result = ReportService().run_quality_check(mapping, data)

        assert result.passed is True
        assert result.injection_failures == []


class _FakeAnthropicUsage:
    def __init__(self) -> None:
        self.input_tokens = 100
        self.output_tokens = 50


class _FakeAnthropicResponse:
    def __init__(self, text: str) -> None:
        self.content = [type("Block", (), {"type": "text", "text": text})()]
        self.usage = _FakeAnthropicUsage()


class TestGenerateFullReport:
    """Full-pipeline orchestration (upload -> parse -> extract -> map ->
    inject) that POST /reports/generate-from-session drives. Only the Claude
    API boundary is faked; uploads-on-disk, parsing, mapping cache, and
    injection all run for real — same approach as
    test_full_pipeline_integration.py, but exercised at the service layer."""

    @pytest.fixture(autouse=True)
    def _fake_claude_api_boundary(self, monkeypatch) -> None:
        monkeypatch.setattr("app.services.claude_service.settings.ANTHROPIC_API_KEY", "sk-test")

        def fake_call_claude(self, prompt, tools, on_status=None):  # noqa: ARG001
            if tools:
                return _FakeAnthropicResponse('{"property_identification": {"district": "Pune"}}')
            return _FakeAnthropicResponse(
                '{"property_identification.district": {"location_id": "t0r0c1", "confidence": "high"},'
                ' "concluded_values.realisable_value_percentage": {"location_id": "t0r1c1", "confidence": "high"},'
                ' "concluded_values.forced_sale_value_percentage": {"location_id": "t0r1c1", "confidence": "high"}}'
            )

        monkeypatch.setattr(ClaudeService, "_call_claude", fake_call_claude)

    def _upload_session(self, tmp_path: Path, monkeypatch) -> str:
        monkeypatch.setattr("app.core.config.settings.UPLOAD_DIR", str(tmp_path / "uploads"))
        session_id = "session-1"

        template_path = tmp_path / "uploads" / session_id / "template" / "template.docx"
        template_path.parent.mkdir(parents=True)
        _build_template(template_path)

        document_path = tmp_path / "uploads" / session_id / "property_document" / "deed.docx"
        document_path.parent.mkdir(parents=True)
        property_document = Document()
        property_document.add_paragraph("Property located in Pune district.")
        property_document.save(document_path)

        return session_id

    def test_full_pipeline_from_uploaded_session(self, tmp_path: Path, monkeypatch) -> None:
        session_id = self._upload_session(tmp_path, monkeypatch)
        output_path = tmp_path / "report.docx"

        injection, quality_check, extracted_data = ReportService().generate_full_report(
            session_id=session_id, output_path=str(output_path)
        )

        assert extracted_data.property_identification.district == "Pune"
        assert "property_identification.district" in injection.filled_fields
        assert quality_check.passed is True
        assert output_path.is_file()

    def test_raises_when_no_template_uploaded(self, tmp_path: Path, monkeypatch) -> None:
        monkeypatch.setattr("app.core.config.settings.UPLOAD_DIR", str(tmp_path / "uploads"))

        with pytest.raises(ReportServiceError, match="template"):
            ReportService().generate_full_report(
                session_id="empty-session", output_path=str(tmp_path / "report.docx")
            )

    def test_raises_when_no_property_documents_uploaded(
        self, tmp_path: Path, monkeypatch
    ) -> None:
        monkeypatch.setattr("app.core.config.settings.UPLOAD_DIR", str(tmp_path / "uploads"))
        session_id = "template-only-session"
        template_path = tmp_path / "uploads" / session_id / "template" / "template.docx"
        template_path.parent.mkdir(parents=True)
        _build_template(template_path)

        with pytest.raises(ReportServiceError, match="property document"):
            ReportService().generate_full_report(
                session_id=session_id, output_path=str(tmp_path / "report.docx")
            )


class TestFinancialYearString:
    def test_january_belongs_to_previous_years_financial_year(self):
        assert _financial_year_string(date(2026, 1, 15)) == "2025-2026"

    def test_april_starts_the_new_financial_year(self):
        assert _financial_year_string(date(2026, 4, 1)) == "2026-2027"

    def test_march_still_belongs_to_the_prior_financial_year(self):
        assert _financial_year_string(date(2027, 3, 31)) == "2026-2027"


class TestCitationFromEasrResult:
    def _result(self, rows, columns) -> EASRGuidelineResult:
        search_input = EASRSearchInput(year="2026-2027", district="Pune", taluka="x", village="y")
        return EASRGuidelineResult(
            search_input=search_input,
            found=bool(rows),
            columns=columns,
            rows=rows,
            source="IGR Maharashtra e ASR (igreval), accessed 2026-08-01",
            accessed_at=datetime(2026, 8, 1, tzinfo=timezone.utc),
        )

    def test_single_row_summarized_without_dropping_data(self):
        row = {
            "Select": "SurveyNo",
            "उपविभाग": "5/52-text",
            "खुली जमीन": "44380",
            "एकक (Rs./)": "चौ. मीटर",
        }
        result = self._result(rows=[row], columns=list(row.keys()))
        citation = _citation_from_easr_result(result)
        assert "Select" not in citation.rate  # excluded — not a rate figure
        assert "44380" in citation.rate
        assert citation.unit == "चौ. मीटर"
        assert citation.access_date == "2026-08-01"

    def test_multiple_rows_are_joined_not_reduced_to_one(self):
        result = self._result(
            rows=[
                {"Assessment Type": "जिरायत", "Rate Rs/-": "100"},
                {"Assessment Type": "बिनशेती", "Rate Rs/-": "200"},
            ],
            columns=["Assessment Type", "Rate Rs/-"],
        )
        citation = _citation_from_easr_result(result)
        assert "100" in citation.rate
        assert "200" in citation.rate
        assert citation.rate.count("|") == 1  # two rows joined by one separator

    def test_empty_rows_never_fabricate_a_rate(self):
        result = self._result(rows=[], columns=[])
        citation = _citation_from_easr_result(result)
        assert citation.rate == INFO_NOT_AVAILABLE


class TestTryAutoEasr:
    """D37: the best-effort auto-lookup step generate_full_report runs after
    extraction when the caller didn't already supply tier1_official_data."""

    def test_skips_when_district_not_extracted(self, monkeypatch):
        data = ValuationReportData()  # district stays at the INFO_NOT_AVAILABLE sentinel
        called = {"count": 0}

        async def fake_try_auto_lookup(*args, **kwargs):  # noqa: ARG001
            called["count"] += 1
            return None

        service = EASRService()
        monkeypatch.setattr(service, "try_auto_lookup", fake_try_auto_lookup)

        ReportService._try_auto_easr(data, date(2026, 8, 1), service)

        assert called["count"] == 0
        assert data.official_rate_evidence == []

    def test_appends_citation_on_confident_match(self, monkeypatch):
        data = ValuationReportData()
        data.property_identification.district = "Mumbai Suburban"
        data.property_identification.taluka = "Borivali"
        data.property_identification.revenue_village = "Kurar"

        fake_result = EASRGuidelineResult(
            search_input=EASRSearchInput(
                year="2026-2027", district="Bombaymains", district_option="Mumbai Suburban",
                taluka="Borivali", village="Kurar (Borivali)",
            ),
            found=True,
            columns=["Select", "उपविभाग", "खुली जमीन"],
            rows=[{"Select": "SurveyNo", "उपविभाग": "72/332-text", "खुली जमीन": "44380"}],
            source="IGR Maharashtra e ASR (igreval), accessed 2026-08-01",
            accessed_at=datetime(2026, 8, 1, tzinfo=timezone.utc),
        )

        async def fake_try_auto_lookup(*args, **kwargs):  # noqa: ARG001
            return fake_result

        service = EASRService()
        monkeypatch.setattr(service, "try_auto_lookup", fake_try_auto_lookup)

        ReportService._try_auto_easr(data, date(2026, 8, 1), service)

        assert len(data.official_rate_evidence) == 1
        assert "44380" in data.official_rate_evidence[0].rate

    def test_leaves_evidence_untouched_when_lookup_returns_none(self, monkeypatch):
        data = ValuationReportData()
        data.property_identification.district = "Some Unknown District"
        data.property_identification.taluka = "x"
        data.property_identification.revenue_village = "y"

        async def fake_try_auto_lookup(*args, **kwargs):  # noqa: ARG001
            return None

        service = EASRService()
        monkeypatch.setattr(service, "try_auto_lookup", fake_try_auto_lookup)

        ReportService._try_auto_easr(data, date(2026, 8, 1), service)

        assert data.official_rate_evidence == []
