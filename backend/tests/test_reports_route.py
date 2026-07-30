import uuid
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.models.template_mapping import FieldMapping
from app.services.claude_service import ClaudeService

client = TestClient(app)


def _build_template(path: Path) -> None:
    document = Document()
    document.add_paragraph("VALUATION REPORT")
    document.add_paragraph("")
    document.save(path)


@pytest.fixture(autouse=True)
def _isolated_storage(tmp_path: Path, monkeypatch) -> None:
    """Points reports storage and the mapping cache at fresh per-test files."""
    monkeypatch.setattr("app.core.config.settings.REPORTS_DIR", str(tmp_path / "reports"))
    monkeypatch.setattr("app.core.db.settings.DATABASE_PATH", str(tmp_path / "test.db"))


@pytest.fixture(autouse=True)
def _fake_claude_mapping(monkeypatch) -> None:
    """No real API call: ClaudeService.map_template_fields is swapped for a fake."""

    def fake_map_template_fields(self, skeleton):  # noqa: ARG001
        return {"property_identification.district": FieldMapping(location_id="p1")}

    monkeypatch.setattr(ClaudeService, "map_template_fields", fake_map_template_fields)


def test_generate_preview_download_flow(tmp_path: Path) -> None:
    template_path = tmp_path / "template.docx"
    _build_template(template_path)

    response = client.post(
        "/reports/generate",
        json={
            "template_path": str(template_path),
            "data": {"property_identification": {"district": "Pune"}},
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["injection"]["filled_fields"] == ["property_identification.district"]
    assert body["quality_check"]["passed"] is True
    report_id = body["report_id"]

    preview = client.get(f"/reports/{report_id}/preview")
    assert preview.status_code == 200
    assert "Pune" in preview.json()["text"]

    download = client.get(f"/reports/{report_id}/download")
    assert download.status_code == 200
    assert download.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def test_preview_unknown_report_id_returns_404() -> None:
    response = client.get(f"/reports/{uuid.uuid4()}/preview")
    assert response.status_code == 404


def test_preview_invalid_report_id_format_returns_400() -> None:
    response = client.get("/reports/not-a-uuid/preview")
    assert response.status_code == 400


def test_generate_with_bad_template_path_returns_400() -> None:
    response = client.post(
        "/reports/generate",
        json={"template_path": "does_not_exist.docx", "data": {}},
    )
    assert response.status_code == 400


class _FakeAnthropicUsage:
    def __init__(self) -> None:
        self.input_tokens = 100
        self.output_tokens = 50


class _FakeAnthropicResponse:
    def __init__(self, text: str) -> None:
        self.content = [type("Block", (), {"type": "text", "text": text})()]
        self.usage = _FakeAnthropicUsage()


class TestGenerateFromSession:
    """POST /reports/generate-from-session: the full upload -> parse ->
    extract -> map -> inject pipeline, driven by session_id alone."""

    @pytest.fixture(autouse=True)
    def _isolated_uploads(self, tmp_path: Path, monkeypatch) -> None:
        monkeypatch.setattr("app.core.config.settings.UPLOAD_DIR", str(tmp_path / "uploads"))
        monkeypatch.setattr("app.services.claude_service.settings.ANTHROPIC_API_KEY", "sk-test")

    def _upload_session(self, tmp_path: Path, session_id: str) -> None:
        template_path = tmp_path / "uploads" / session_id / "template" / "template.docx"
        template_path.parent.mkdir(parents=True)
        _build_template(template_path)

        document_path = tmp_path / "uploads" / session_id / "property_document" / "deed.docx"
        document_path.parent.mkdir(parents=True)
        property_document = Document()
        property_document.add_paragraph("Property located in Pune district.")
        property_document.save(document_path)

    def test_full_pipeline_from_session_id(self, tmp_path: Path, monkeypatch) -> None:
        session_id = "session-route-1"
        self._upload_session(tmp_path, session_id)

        def fake_call_claude(self, prompt, tools, on_status=None):  # noqa: ARG001
            if tools:
                return _FakeAnthropicResponse('{"property_identification": {"district": "Pune"}}')
            return _FakeAnthropicResponse(
                '{"property_identification.district": {"location_id": "t0r0c1", "confidence": "high"}}'
            )

        monkeypatch.setattr(ClaudeService, "_call_claude", fake_call_claude)

        response = client.post(
            "/reports/generate-from-session", json={"session_id": session_id}
        )

        assert response.status_code == 200
        body = response.json()
        assert body["extracted_data"]["property_identification"]["district"] == "Pune"
        assert body["quality_check"]["passed"] is True

        preview = client.get(f"/reports/{body['report_id']}/preview")
        assert "Pune" in preview.json()["text"]

        progress = client.get(f"/reports/progress/{session_id}")
        assert progress.json() == {"percent": 100, "stage": "Done"}

    def test_progress_defaults_to_zero_for_an_unknown_session(self) -> None:
        response = client.get("/reports/progress/never-started")
        assert response.json() == {"percent": 0, "stage": "idle"}

    def test_no_uploaded_documents_returns_400(self) -> None:
        response = client.post(
            "/reports/generate-from-session", json={"session_id": "nonexistent-session"}
        )
        assert response.status_code == 400
